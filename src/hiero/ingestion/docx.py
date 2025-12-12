from __future__ import annotations

from hashlib import sha256
from typing import BinaryIO

from docx import Document as DocxDocument

from .base import Document, DocumentMetadata, DocumentSource, IngestionError, IngestorProtocol


class DOCXIngestor(IngestorProtocol):
    async def ingest(
        self,
        source: BinaryIO,
        metadata: DocumentMetadata | None = None,
    ) -> Document:
        try:
            source.seek(0)
            docx = DocxDocument(source)
        except Exception as exc:
            raise IngestionError(f"Failed to open DOCX: {exc}") from exc

        text_parts = [p.text for p in docx.paragraphs if p.text and p.text.strip()]
        content = "\n".join(text_parts)
        content_hash = sha256(content.encode("utf-8")).hexdigest()

        core = getattr(docx, "core_properties", None)
        extracted = DocumentMetadata(
            title=getattr(core, "title", None) or (metadata.title if metadata else None),
            author=getattr(core, "author", None) or (metadata.author if metadata else None),
            file_type="docx",
            **(metadata.model_dump(exclude_unset=True) if metadata else {}),
        )

        return Document(
            content=content,
            metadata=extracted,
            source=DocumentSource.FILE_UPLOAD,
            content_hash=content_hash,
        )

    def supports(self, file_type: str) -> bool:
        return file_type.lower() == "docx"

